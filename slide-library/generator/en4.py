# Delivery chapters for customer projects: kick-off, analysis, target picture, validation, implementation, workshops, appendix
from en import *
from kh import _strip_style
from pptx.enum.dml import MSO_LINE

RAG_EN = ('On track', 'At risk', 'Critical', 'Not started')


def status_col(s, tb_y, rowh, nrows, x, states, labels=('Done', 'In progress', 'Overdue', 'Open')):
    m = dict(zip('gyrn', labels))
    for r in range(nrows):
        y = tb_y + (r + 1) * rowh
        rag(s, x, y + rowh / 2 - 0.1, states[r])
        text(s, x + 0.3, y, 1.6, rowh, m[states[r]], size=13, anchor='m')


def tag(s, x, y, t, fill=BLUE, w=None):
    w = w or 0.25 + 0.09 * len(t)
    return rrect(s, x, y, w, 0.3, fill=fill, radius=0.15, content=t, size=11, color=WHITE if fill in (BLUE, BLACK, TAUPE) else BLACK,
                 align='c', anchor='m', margin=0)


# =====================================================================================
def kickoff(D):
    C = 'Kick-off'
    divider(D, 0, 'Kick-off &\nways of\nworking', N('Everything needed to start a customer project: objectives, team, rules, cadence, communication, data and interviews.',
                                                    'Use in the kick-off meeting and keep updated in the project handbook.',
                                                    'A good kick-off saves weeks later: agree objectives, roles and rules explicitly.'))
    s = cs(D, C, 'Objectives and success criteria', 'ko_objectives',
           N('Makes success measurable from day one.', 'Three objectives, each with KPI, baseline and target.',
             'Agree the success criteria with the sponsor – they are the basis for the closing report and the reference.'))
    rows = [['Objective', 'KPI', 'Baseline', 'Target', 'By when']] + [['Objective %d in one line' % i, '[KPI]', '[0]', '[0]', '[MM/YYYY]'] for i in range(1, 4)]
    tb = table(s, L, 1.5, W, rows, colw=[4.6, 2.6, 1.6, 1.6, 1.8], rowh=[0.45, 0.85, 0.85, 0.85], size=15, aligns=['l', 'l', 'c', 'c', 'c'])
    for r in range(1, 4):
        fill_tf(tb.cell(r, 3).text_frame, '[0]', size=16, color=BLUE, align='c')
    sowhat(s, KT)

    s = cs(D, C, 'Team and contacts', 'ko_contacts', N('Contact list of the joint project team.', 'Add rows as needed; keep in sync with the project handbook.'),
           title='Project team and contacts')
    rows = [['Role', 'Name', 'Organisation', 'E-mail', 'Phone', 'Availability']] + [[r, '[Name]', o, '[e-mail]', '[phone]', '[% / days]'] for r, o in
                                                                                     [('Sponsor', '[Customer]'), ('Project lead', '[Customer]'), ('Project lead', 'Körber'),
                                                                                      ('Workstream lead', '[Customer]'), ('Workstream lead', 'Körber'), ('Expert', 'Körber')]]
    table(s, L, 1.5, W, rows, colw=[2.2, 2.2, 1.9, 2.8, 1.8, 1.7], rowh=0.6, size=14)

    s = cs(D, C, 'Ways of working', 'ko_rules', N('Collaboration rules agreed in the kick-off.', 'Agree the rules with the team, not for the team.'),
           title='Six rules for working together')
    rules = [('handshake', 'One team', 'Joint teams, shared goals, no “us and them”.'), ('schedule', 'Fast decisions', 'Decisions within [x] working days.'),
             ('visibility', 'Transparency', 'Risks and issues are raised early and openly.'), ('folder', 'Single source of truth', 'All documents in [SharePoint / Teams].'),
             ('event', 'Fixed rhythm', 'Meetings follow the agreed calendar.'), ('school', 'Knowledge transfer', 'Your team is enabled to run the solution.')]
    cw = (W - 2 * 0.5) / 3
    for i, (ic, h, b) in enumerate(rules):
        x = L + (i % 3) * (cw + 0.5); y = 1.5 + (i // 3) * 2.15
        icon_circle(s, ic, x, y, 0.62)
        text(s, x + 0.85, y + 0.12, cw - 0.85, 0.4, h, size=18)
        hline(s, x, y + 0.85, cw, STONE, 0.5)
        text(s, x, y + 1.0, cw, 0.8, b, size=14)

    s = cs(D, C, 'Meeting calendar', 'ko_calendar', N('Weekly meeting rhythm of the project.', 'Move and rename the meeting blocks.',
                                                       'Fix the rhythm for the full project in week 1 – calendars of experts fill up quickly.'),
           title='Weekly meeting rhythm')
    days = ['Monday', 'Tuesday', 'Wednesday', 'Thursday', 'Friday']
    x0 = L + 1.2; dw = (R - x0) / 5
    for i, d in enumerate(days):
        text(s, x0 + i * dw, 1.45, dw, 0.35, d, size=15, color=BLUE, align='c')
    hline(s, L, 1.85, W, BLACK)
    hours = ['09:00', '11:00', '13:00', '15:00', '17:00']
    for j, h in enumerate(hours):
        y = 1.95 + j * 0.8
        text(s, L, y, 1.1, 0.3, h, size=12, color=DGRAY)
        hline(s, x0, y, R - x0, LSAND, 0.5)
    meets = [(0, 0.0, 0.6, 'Project leads', BLUE), (2, 0.0, 0.6, 'Project leads', BLUE), (1, 1.6, 1.4, 'Workstream\nmeeting', SAND),
             (3, 1.6, 1.4, 'Workstream\nmeeting', SAND), (4, 0.0, 1.0, 'Steering committee\n(monthly)', BLACK), (4, 2.4, 0.8, 'Weekly status', MBLUE),
             (0, 2.4, 1.0, 'Workshop', SAND)]
    for d, y0, h, t, f in meets:
        rect(s, x0 + d * dw + 0.06, 1.97 + y0, dw - 0.12, h, fill=f, content=t, size=12, color=WHITE if f in (BLUE, BLACK) else BLACK, anchor='m', align='c')

    s = cs(D, C, 'Communication plan', 'ko_comms', N('Who gets which information, when and how.', 'Adapt audiences and channels to the customer organisation.'))
    rows = [['Audience', 'Information', 'Channel', 'Frequency', 'Owner']] + [[a, '[Content]', c, f, '[Name]'] for a, c, f in
                                                                          [('Steering committee', 'Meeting + report', 'Monthly'), ('Project team', 'Teams / jour fixe', 'Weekly'),
                                                                           ('Affected employees', 'Town hall, intranet', 'Milestones'), ('Management', 'E-mail update', 'Monthly'),
                                                                           ('Works council', 'Meeting', 'As required')]]
    table(s, L, 1.5, W, rows, colw=[2.8, 3.4, 2.6, 1.8, 2.0], rowh=0.65, size=15)

    s = cs(D, C, 'Information request', 'ko_datarequest', N('List of documents and data needed from the customer.', 'Send before the kick-off; track weekly.',
                                                            'Missing data is the most common cause of delay in the analysis phase.'),
           title='Information request: status of requested documents and data', source='As of: [date]')
    rows = [['#', 'Document / data', 'Owner customer', 'Due', 'Status']] + [[str(i), '[Document or data set]', '[Name]', '[DD.MM.]', ''] for i in range(1, 7)]
    table(s, L, 1.5, W, rows, colw=[0.6, 5.8, 2.6, 1.6, 2.0], rowh=0.58, size=14)
    status_col(s, 1.5, 0.58, 6, L + W * (10.6 / 12.6) + 0.15, list('ggyrnn'), labels=('Received', 'Partly', 'Overdue', 'Open'))

    s = cs(D, C, 'Interview plan', 'ko_interviews', N('Plan of stakeholder interviews in the analysis phase.', 'One row per interview; 8–12 interviews are typical.',
                                                       'Send the interview guide in advance – answers get more precise.'), source='As of: [date]')
    rows = [['Interviewee', 'Function', 'Focus topics', 'Date', 'Körber', 'Status']] + [['[Name]', '[Function]', '[Topics]', '[DD.MM.]', '[Name]', ''] for _ in range(6)]
    table(s, L, 1.5, W, rows, colw=[2.2, 2.2, 3.6, 1.4, 1.6, 1.6], rowh=0.58, size=14)
    status_col(s, 1.5, 0.58, 6, L + W * (11.0 / 12.6) + 0.15, list('ggynnn'), labels=('Done', 'Scheduled', 'Overdue', 'Open'))


# =====================================================================================
def analysis(D):
    C = 'Analysis'
    divider(D, 0, 'Analysis &\ndiagnostics', N('Slides for the analysis readout: findings, assessment, gaps, survey, site visit.',
                                               'Every finding needs evidence and an implication.', 'Findings → implications → recommendations: never stop at the finding.'))
    s = cs(D, C, 'Key findings', 'an_findings', N('Overview of the main findings of the analysis.', 'Maximum five findings; each with evidence and implication.'))
    rows = [['#', 'Finding', 'Evidence', 'Implication']] + [[str(i), 'Finding in one sentence', '[Data point, source]', 'What it means for the customer'] for i in range(1, 6)]
    tb = table(s, L, 1.5, W, rows, colw=[0.6, 4.4, 3.4, 4.2], rowh=[0.45] + [0.72] * 5, size=14)
    for r in range(1, 6):
        fill_tf(tb.cell(r, 0).text_frame, str(r), size=18, color=BLUE)

    s = cs(D, C, 'Assessment heatmap', 'an_heatmap', N('Assessment of several units (sites, lines, departments) across dimensions.',
                                                       'Score 1–4; recolour the cells: dark blue = strong, orange = weak.'), source='Source: [Assessment, date]')
    dims = ['Dimension %d' % i for i in range(1, 7)]
    units = ['Unit %d' % i for i in range(1, 6)]
    sc = [[4, 3, 2, 3, 1, 3], [3, 3, 3, 2, 2, 4], [2, 1, 2, 3, 1, 2], [4, 4, 3, 3, 3, 3], [3, 2, 1, 2, 2, 3]]
    col = {4: BLUE, 3: MBLUE, 2: SAND, 1: ORANGE}
    x0 = L + 2.2; cw = (R - x0 - 1.4) / 6; rh = 0.68
    for j, d in enumerate(dims):
        text(s, x0 + j * cw, 1.45, cw, 0.4, d, size=13, align='c', color=BLUE)
    text(s, R - 1.3, 1.45, 1.3, 0.4, 'Average', size=13, align='c', color=BLUE)
    for i, u in enumerate(units):
        y = 1.95 + i * rh
        text(s, L, y, 2.1, rh - 0.08, u, size=15, anchor='m')
        for j in range(6):
            v = sc[i][j]
            rect(s, x0 + j * cw + 0.04, y, cw - 0.08, rh - 0.08, fill=col[v], content=str(v), size=15, color=WHITE if v in (4, 1) else BLACK, align='c', anchor='m')
        text(s, R - 1.3, y, 1.3, rh - 0.08, '%.1f' % (sum(sc[i]) / 6), size=16, align='c', anchor='m')
    for k, (v, t) in enumerate([(4, 'Strong'), (3, 'Good'), (2, 'Partial'), (1, 'Weak')]):
        rect(s, L + k * 1.6, 5.55, 0.22, 0.22, fill=col[v]); text(s, L + k * 1.6 + 0.3, 5.51, 1.2, 0.3, '%d  %s' % (v, t), size=12)

    s = cs(D, C, 'Gap analysis', 'an_gaps', N('Current state vs. target state with the resulting measures.', 'One row per dimension; gap in blue.'))
    rows = [['Dimension', 'Current state', 'Target state', 'Gap', 'Measure']] + [['Dimension %d' % i, '[Today]', '[Target]', '[Gap]', '[Measure]'] for i in range(1, 6)]
    tb = table(s, L, 1.5, W, rows, colw=[2.2, 2.8, 2.8, 1.6, 3.2], rowh=[0.45] + [0.7] * 5, size=14)
    for r in range(1, 6):
        c = tb.cell(r, 3); c.fill.solid(); setc(c.fill.fore_color, LBLUE); fill_tf(c.text_frame, '[Gap]', size=14, color=BLUE, anchor='m')

    s = cs(D, C, 'Survey results', 'an_survey', N('Results of an employee or customer survey (Likert scale).', 'Edit chart data; sort statements by agreement.'),
           source='Source: [Survey], n = [xx], [date]')
    stm = ['Statement %d' % i for i in range(1, 7)]
    bar(s, L, 1.4, 9.0, 4.45, stm, [('Fully agree', [0.45, 0.38, 0.30, 0.22, 0.15, 0.10]), ('Agree', [0.30, 0.32, 0.30, 0.28, 0.25, 0.20]),
                                    ('Neutral', [0.15, 0.18, 0.20, 0.25, 0.25, 0.25]), ('Disagree', [0.10, 0.12, 0.20, 0.25, 0.35, 0.45])],
        kind='bar', stacked=True, pct=True, colors=[BLUE, MBLUE, SAND, STONE], legend='b', fmt='0%', gap=45, size=12)
    kpi(s, 9.9, 1.6, R - 9.9, 'n = 000', 'Participants', '00% response rate', bigsize=28)
    kpi(s, 9.9, 3.2, R - 9.9, '00%', 'Average agreement', bigsize=28)
    sowhat(s, KT)

    s = cs(D, C, 'Site visit observations', 'an_gemba', N('Observations from a site visit or Gemba walk.', 'Replace the picture; keep positive and improvement observations separate.'))
    picture(s, IMG + 'image21.jpeg', L, 1.5, 4.2, 4.3)
    cw = (R - (L + 4.6) - 0.3) / 2
    for i, (h, ic, c) in enumerate([('What works well', 'thumb_up', BLUE), ('Improvement potential', 'build', ORANGE)]):
        x = L + 4.6 + i * (cw + 0.3)
        icon(s, ic, x, 1.5, 0.42, c)
        text(s, x + 0.6, 1.52, cw - 0.6, 0.4, h, size=18)
        hline(s, x, 2.1, cw, BLACK)
        for k in range(3):
            y = 2.25 + k * 1.15
            text(s, x, y, cw, 0.35, 'Observation', size=15)
            text(s, x, y + 0.38, cw, 0.6, '[Where, what, impact]', size=13, color=DGRAY)

    s = cs(D, C, 'Fact base', 'an_facts', N('Key facts and figures about the customer situation.', 'Six figures, each with source.'), source='Source: [Source]')
    cw = (W - 2 * 0.5) / 3
    for i in range(6):
        x = L + (i % 3) * (cw + 0.5); y = 1.5 + (i // 3) * 2.0
        kpi(s, x, y, cw, ['00%', '€0.0m', '000', '0.0 d', '00', '0x'][i], 'Label of the figure', '[Source / context]', bigsize=36)
    sowhat(s, KT)


# =====================================================================================
def target(D):
    C = 'Target picture'
    divider(D, 0, 'Target\npicture &\nroadmap', N('From design principles to the target operating model, requirements, options and roadmap.',
                                                 'Fill with the agreed concept; update after each design workshop.',
                                                 'The roadmap is the bridge from concept to implementation – make it the most concrete slide.'))
    s = cs(D, C, 'Design principles', 'tp_principles', N('Guiding principles for the target design – agreed with the sponsor.', 'Five principles; each with a “this means” line.'))
    for i in range(5):
        y = 1.5 + i * 0.85
        num(s, L, y + 0.1, i + 1, 0.5)
        text(s, L + 0.8, y + 0.12, 4.0, 0.45, 'Principle %d' % (i + 1), size=18)
        text(s, L + 5.0, y + 0.12, 7.6, 0.55, 'This means: [what it implies for the design]', size=15)
        if i < 4:
            hline(s, L + 0.8, y + 0.8, W - 0.8, STONE, 0.5)

    s = cs(D, C, 'Target operating model', 'tp_tom', N('House of the target operating model.', 'Adapt pillars to the scope (e.g. processes, organisation, systems, data).'))
    shape(s, SH.ISOSCELES_TRIANGLE, L + 1.5, 1.45, W - 3.0, 0.9, fill=BLUE)
    text(s, L + 1.5, 1.8, W - 3.0, 0.5, 'Vision / target picture in one line', size=15, color=WHITE, align='c', anchor='m')
    pw = (W - 3.0 - 3 * 0.2) / 4
    for i, (h, ic) in enumerate([('Processes', 'account_tree'), ('Organisation', 'groups'), ('Systems & data', 'database'), ('Governance', 'gavel' if False else 'verified')]):
        x = L + 1.5 + i * (pw + 0.2)
        rect(s, x, 2.5, pw, 2.4, fill=LSAND)
        icon(s, ic, x + 0.2, 2.65, 0.4, BLUE)
        text(s, x + 0.2, 3.15, pw - 0.4, 0.4, h, size=16)
        text(s, x + 0.2, 3.6, pw - 0.4, 1.2, bl(['Element', 'Element', 'Element'], sa=4), size=13)
    rect(s, L + 1.5, 5.05, W - 3.0, 0.65, fill=SAND, content='Foundation: GxP compliance, culture and capabilities', size=15, align='c', anchor='m')

    s = cs(D, C, 'Requirements (MoSCoW)', 'tp_requirements', N('Prioritised requirements, e.g. for a URS or system selection.', 'Priority: Must / Should / Could / Won’t.',
                                                                'Maximum 60% “Must” – otherwise prioritisation has not really happened.'))
    rows = [['ID', 'Requirement', 'Category', 'Priority', 'Source']] + [['REQ-%02d' % i, '[Requirement in one sentence]', '[Category]', '', '[Stakeholder]'] for i in range(1, 7)]
    table(s, L, 1.5, W, rows, colw=[1.2, 5.6, 2.0, 1.6, 2.2], rowh=0.58, size=14)
    prio = ['Must', 'Must', 'Should', 'Should', 'Could', 'Won’t']
    pc = {'Must': BLUE, 'Should': MBLUE, 'Could': SAND, 'Won’t': STONE}
    xp = L + W * (8.8 / 12.6) + 0.1
    for r, p in enumerate(prio):
        tag(s, xp, 1.5 + (r + 1) * 0.58 + 0.14, p, fill=pc[p], w=1.2)

    s = cs(D, C, 'Long list to short list', 'tp_funnel', N('Transparent narrowing of options or vendors.', 'Name the screening criteria per stage.'))
    stages = [('Long list', '00 options', 'Desk research, market scan'), ('Screening', '00 options', 'Knock-out criteria'),
              ('Short list', '0 options', 'Detailed evaluation, demos'), ('Recommendation', '1 option', 'Business case, references')]
    for i, (h, n_, crit) in enumerate(stages):
        w_ = 6.4 - i * 1.2; y = 1.5 + i * 1.08
        x = L + (6.4 - w_) / 2
        rect(s, x, y, w_, 0.9, fill=BLUE if i == 3 else (MBLUE if i == 2 else SAND), content=[{'t': h, 'size': 15}, {'t': n_, 'size': 13}],
             color=WHITE if i == 3 else BLACK, align='c', anchor='m')
        text(s, 7.3, y, R - 7.3, 0.9, [{'t': [('Criteria:  ', {'color': BLUE}), (crit, {})]}], size=15, anchor='m')
        hline(s, 7.3, y + 0.98, R - 7.3, STONE, 0.5)

    s = cs(D, C, 'Vendor / system evaluation', 'tp_vendor', N('Weighted evaluation of vendors or systems.', 'Scores 1–5; weighted total = Σ weight × score.'))
    rows = [['Criterion', 'Weight', 'Vendor A', 'Vendor B', 'Vendor C']] + [['Criterion %d' % i, '00%', '0', '0', '0'] for i in range(1, 6)] + \
           [[('Weighted total', {}), '100%', ('0.0', {'size': 17}), ('0.0', {'size': 17, 'color': WHITE, 'fill': BLUE}), ('0.0', {'size': 17})]]
    table(s, L, 1.5, W, rows, colw=[4.0, 1.6, 2.3, 2.3, 2.3], rowh=0.58, size=15, aligns=['l', 'c', 'c', 'c', 'c'],
          fills={(r, 3): LBLUE for r in range(1, 6)})
    sowhat(s, KT)

    s = cs(D, C, 'Initiative portfolio', 'tp_portfolio', N('Overview of all initiatives with value, effort and wave.', 'Harvey balls 0–4; sort by wave.'))
    rows = [['Initiative', 'Owner', 'Value', 'Effort', 'Wave', 'Benefit p.a.']] + [['Initiative %d' % i, '[Name]', '', '', '', '€0.0m'] for i in range(1, 7)]
    table(s, L, 1.5, W, rows, colw=[4.0, 2.0, 1.4, 1.4, 1.6, 2.2], rowh=0.58, size=14, aligns=['l', 'l', 'c', 'c', 'c', 'r'])
    xs = [L + W * sum([4.0, 2.0, 1.4, 1.4, 1.6, 2.2][:k]) / 12.6 for k in range(7)]
    vals = [(4, 1, 1), (3, 2, 1), (4, 3, 2), (2, 1, 1), (3, 4, 3), (2, 3, 3)]
    for r, (v, e, w_) in enumerate(vals):
        y = 1.5 + (r + 1) * 0.58 + 0.11
        harvey(s, (xs[2] + xs[3]) / 2 - 0.18, y, 0.36, v)
        harvey(s, (xs[3] + xs[4]) / 2 - 0.18, y, 0.36, e)
        tag(s, (xs[4] + xs[5]) / 2 - 0.45, y + 0.03, 'Wave %d' % w_, fill=BLUE if w_ == 1 else SAND, w=0.9)

    s = cs(D, C, 'Initiative charter', 'tp_charter', N('One-pager per initiative – basis for prioritisation and steering.', 'Duplicate per initiative.'),
           title='Initiative [name]: [objective in one line]')
    left = [('Objective', '[What should be achieved?]'), ('Scope', '[In / out of scope]'), ('Approach', '[Main steps]')]
    for i, (h, b) in enumerate(left):
        y = 1.5 + i * 1.1
        hline(s, L, y, 6.6, BLACK); label(s, L, y + 0.12, 6.6, h); text(s, L, y + 0.42, 6.6, 0.6, b, size=15)
    xr = 7.4; wr = R - xr
    rect(s, xr, 1.5, wr, 3.2, fill=LSAND)
    for i, (h, b) in enumerate([('Owner', '[Name]'), ('Timeline', '[MM/YY – MM/YY]'), ('Budget', '€[0]k'), ('KPI / benefit', '[KPI, target]'), ('Dependencies', '[Initiatives]')]):
        y = 1.65 + i * 0.6
        text(s, xr + 0.25, y, 1.9, 0.5, h, size=13, color=BLUE, anchor='m'); text(s, xr + 2.2, y, wr - 2.4, 0.5, b, size=14, anchor='m')
    label(s, xr, 4.95, wr, 'Key risks')
    text(s, xr, 5.25, wr, 0.6, '[Risk and mitigation]', size=14)

    s = cs(D, C, 'Roadmap in waves', 'tp_roadmap', N('Implementation roadmap by workstream and wave.', 'Move bars; quick wins in wave 1.',
                                                     'Wave 1 must deliver visible results within six months – that secures sponsorship.'))
    waves = [('Wave 1: Quick wins', '0–6 months'), ('Wave 2: Build', '6–18 months'), ('Wave 3: Scale', '18–36 months')]
    x0 = L + 2.4; ww = (R - x0) / 3
    for i, (h, t) in enumerate(waves):
        rect(s, x0 + i * ww + 0.04, 1.45, ww - 0.08, 0.7, fill=BLUE if i == 0 else SAND, content=[{'t': h, 'size': 14}, {'t': t, 'size': 11}],
             color=WHITE if i == 0 else BLACK, align='c', anchor='m')
    bars = [(0, 0.0, 0.9), (0, 1.0, 2.0), (1, 0.2, 1.4), (1, 1.5, 2.9), (2, 0.0, 0.8), (2, 1.0, 2.6), (3, 0.5, 2.0), (3, 2.1, 3.0)]
    for j in range(4):
        y = 2.35 + j * 0.85
        text(s, L, y, 2.3, 0.7, 'Workstream %d' % (j + 1), size=15, anchor='m')
        hline(s, L, y + 0.78, W, LSAND, 0.75)
    for k, (j, a, b_) in enumerate(bars):
        y = 2.35 + j * 0.85 + 0.12
        rect(s, x0 + a * ww + 0.04, y, (b_ - a) * ww - 0.08, 0.46, fill=BLUE if a < 1 else STONE, content='Initiative', size=12,
             color=WHITE if a < 1 else BLACK, anchor='m', margin=(0.1, 0, 0.05, 0))
    sowhat(s, KT)

    s = cs(D, C, 'Quick wins', 'tp_quickwins', N('Measures that can be implemented within weeks – build momentum.', 'Only measures with owner and date.'))
    rows = [['#', 'Quick win', 'Benefit', 'Effort', 'Owner', 'Date']] + [[str(i), '[Measure]', '[Benefit]', '', '[Name]', '[DD.MM.]'] for i in range(1, 7)]
    table(s, L, 1.5, W, rows, colw=[0.6, 4.6, 3.0, 1.2, 1.8, 1.4], rowh=0.58, size=14, aligns=['l', 'l', 'l', 'c', 'l', 'l'])
    xe = L + W * (8.2 + 0.6) / 12.6
    for r, e in enumerate([1, 1, 2, 1, 2, 1]):
        harvey(s, xe - 0.18, 1.5 + (r + 1) * 0.58 + 0.11, 0.36, e)


# =====================================================================================
def validation(D):
    C = 'Validation & compliance'
    divider(D, 0, 'Validation\n& compliance', N('GxP-specific slides: V-model, deliverables, risk assessment, traceability, qualification status, inspection readiness.',
                                               'Align terms with the customer’s quality system (SOPs, templates).',
                                               'This is where Körber Pharma Consulting differentiates – show regulatory depth.'))
    s = cs(D, C, 'V-model', 'va_vmodel', N('Validation approach following the V-model (GAMP 5).', 'Adapt the levels to the system category and the customer’s SOPs.'),
           title='Validation approach: specification and verification follow the V-model')
    left = [('URS', 'User requirements'), ('FS', 'Functional specification'), ('DS', 'Design specification')]
    right = [('PQ', 'Performance qualification'), ('OQ', 'Operational qualification'), ('IQ', 'Installation qualification')]
    bw, bh = 3.0, 0.8
    for i, (a, b) in enumerate(left):
        x = L + i * 1.2; y = 1.5 + i * 1.1
        rect(s, x, y, bw, bh, fill=SAND, content=[{'t': a, 'size': 16, 'color': BLUE}, {'t': b, 'size': 12}], anchor='m', margin=(0.2, 0.05, 0.1, 0.05))
        xr_ = R - bw - i * 1.2
        rect(s, xr_, y, bw, bh, fill=LBLUE, content=[{'t': right[i][0], 'size': 16, 'color': BLUE}, {'t': right[i][1], 'size': 12}], anchor='m', margin=(0.2, 0.05, 0.1, 0.05))
        arrow(s, xr_, y + bh / 2, x + bw, y + bh / 2, color=STONE, lw=1)
        text(s, x + bw + 0.2, y + bh / 2 - 0.32, xr_ - x - bw - 0.4, 0.3, 'verifies', size=11, align='c', color=DGRAY)
    rect(s, L + 3.6, 4.85, W - 7.2, 0.8, fill=BLUE, content='Build / configuration', size=16, color=WHITE, align='c', anchor='m')
    label(s, L, 5.0, 3, 'Specification ↓', color=DGRAY); label(s, R - 3, 5.0, 3, 'Verification ↑', color=DGRAY, align='r')

    s = cs(D, C, 'Validation deliverables', 'va_deliverables', N('Who creates, reviews and approves which validation document.', 'C = create, R = review, A = approve.'))
    roles = ['Körber', 'System owner', 'QA', 'Vendor']
    docs = [('Validation plan', 'C R A -'), ('Risk assessment', 'C R A C'), ('URS', 'C A R -'), ('IQ / OQ protocols', 'C R A C'),
            ('PQ protocol', 'C A A -'), ('Traceability matrix', 'C R A -'), ('Validation report', 'C R A -')]
    data = [['Deliverable'] + roles] + [[d] + p.split() for d, p in docs]
    tb = table(s, L, 1.5, W, data, colw=[4.4] + [2.05] * 4, rowh=0.55, size=15, aligns=['l'] + ['c'] * 4)
    for r in range(1, len(data)):
        for c in range(1, 5):
            v = data[r][c]
            if v == 'A':
                cell = tb.cell(r, c); cell.fill.solid(); setc(cell.fill.fore_color, BLUE); fill_tf(cell.text_frame, 'A', size=15, color=WHITE, align='c')
            elif v == 'C':
                cell = tb.cell(r, c); cell.fill.solid(); setc(cell.fill.fore_color, LBLUE); fill_tf(cell.text_frame, 'C', size=15, color=BLUE, align='c')
    text(s, L, 6.0, W, 0.3, [{'t': [('C ', {'color': BLUE}), ('Create    ', {}), ('R ', {'color': BLUE}), ('Review    ', {}), ('A ', {'color': BLUE}), ('Approve', {})]}], size=12)

    s = cs(D, C, 'Risk assessment (FMEA)', 'va_fmea', N('Risk-based approach: FMEA for functions or process steps.', 'Severity, occurrence, detection 1–5; RPN = S × O × D.',
                                                       'Agree the RPN thresholds with QA before the assessment.'))
    rows = [['Function / step', 'Failure mode', 'Effect', 'S', 'O', 'D', 'RPN', 'Mitigation']] + \
           [['[Function]', '[Failure mode]', '[Effect]', s_, o, d, str(int(s_) * int(o) * int(d)), '[Measure]'] for s_, o, d in
            [('5', '3', '4'), ('4', '3', '3'), ('3', '2', '4'), ('4', '2', '2'), ('2', '2', '2')]]
    tb = table(s, L, 1.5, W, rows, colw=[2.2, 2.2, 2.0, 0.6, 0.6, 0.6, 0.9, 3.5], rowh=0.62, size=14, aligns=['l', 'l', 'l', 'c', 'c', 'c', 'c', 'l'])
    for r in range(1, 6):
        v = int(rows[r][6]); c = tb.cell(r, 6); c.fill.solid(); setc(c.fill.fore_color, ORANGE if v >= 36 else (AMBER if v >= 20 else SAND))
        fill_tf(c.text_frame, str(v), size=14, align='c')
    for k, (f, t) in enumerate([(ORANGE, 'RPN ≥ 36: mitigation required'), (AMBER, 'RPN 20–35: mitigation recommended'), (SAND, 'RPN < 20: acceptable')]):
        rect(s, L + k * 4.2, 5.5, 0.22, 0.22, fill=f); text(s, L + k * 4.2 + 0.3, 5.46, 3.8, 0.3, t, size=12)

    s = cs(D, C, 'Traceability matrix', 'va_trace', N('Proof that every requirement is tested.', 'One row per requirement; ✓ where a test covers it.'))
    tests = ['IQ-01', 'IQ-02', 'OQ-01', 'OQ-02', 'OQ-03', 'PQ-01']
    data = [['Requirement', 'Risk'] + tests] + [['REQ-%02d [short text]' % i, ['High', 'High', 'Medium', 'Medium', 'Low', 'Low'][i - 1]] + [''] * 6 for i in range(1, 7)]
    tb = table(s, L, 1.5, W, data, colw=[3.8, 1.4] + [1.23] * 6, rowh=0.58, size=14, aligns=['l', 'c'] + ['c'] * 6)
    cover = [[0, 2], [2, 3], [1, 4], [3], [5], [4, 5]]
    xs = [L + W * sum(([3.8, 1.4] + [1.23] * 6)[:k]) / (3.8 + 1.4 + 1.23 * 6) for k in range(9)]
    for r, cols in enumerate(cover):
        for c in cols:
            icon(s, 'check_circle', (xs[2 + c] + xs[3 + c]) / 2 - 0.15, 1.5 + (r + 1) * 0.58 + 0.14, 0.3, BLUE)
    sowhat(s, KT, lead='Coverage')

    s = cs(D, C, 'Qualification status', 'va_status', N('Status of commissioning and qualification per system or equipment.', 'Recolour dots; update weekly during CQV.'),
           source='As of: [date]')
    phases = ['FAT', 'SAT', 'IQ', 'OQ', 'PQ', 'Release']
    data = [['System / equipment'] + phases] + [['System %d' % i] + [''] * 6 for i in range(1, 7)]
    table(s, L, 1.5, W, data, colw=[3.6] + [1.5] * 6, rowh=0.58, size=14, aligns=['l'] + ['c'] * 6)
    st = ['gggggy', 'ggggyn', 'gggynn', 'ggyrnn', 'gynnnn', 'ynnnnn']
    xs = [L + W * sum(([3.6] + [1.5] * 6)[:k]) / 12.6 for k in range(8)]
    for r in range(6):
        for c in range(6):
            rag(s, (xs[1 + c] + xs[2 + c]) / 2 - 0.11, 1.5 + (r + 1) * 0.58 + 0.18, st[r][c], 0.22)
    legend_rag(s, L, 5.75, size=11, labels=('Completed', 'In progress', 'Deviation open', 'Not started'))

    s = cs(D, C, 'Inspection readiness', 'va_readiness', N('Readiness for an authority inspection or customer audit by area.', 'Bar length = readiness in %; orange below 70%.'),
           source='Source: [Self-assessment / mock inspection], [date]')
    areas = ['Quality system', 'Documentation', 'Data integrity', 'Facilities & equipment', 'Contamination control', 'Training']
    vals = [0.9, 0.82, 0.62, 0.85, 0.68, 0.95]
    bar(s, L, 1.4, 8.6, 4.45, areas, [('Readiness', vals)], kind='bar', fmt='0%', gap=40, size=13,
        point_colors=[ORANGE if v < 0.7 else BLUE for v in vals], max_=1)
    kpi(s, 9.5, 1.6, R - 9.5, '00%', 'Overall readiness', bigsize=32)
    kpi(s, 9.5, 3.2, R - 9.5, '0', 'Critical gaps open', bigsize=32)
    sowhat(s, KT)


# =====================================================================================
def implementation(D):
    C = 'Implementation'
    divider(D, 0, 'Implemen-\ntation &\nchange', N('Running the implementation: open points, decisions, resources, change, training, go-live, cut-over.',
                                                 'Update weekly; keep the same format throughout the project.',
                                                 'Projects fail on adoption, not on technology – plan change management from the start.'))
    s = cs(D, C, 'Open points list', 'im_opl', N('Action items / open points list – the working tool of the project team.', 'One row per action; review weekly.'),
           title='Open points: status of actions', source='As of: [date]')
    rows = [['#', 'Action', 'Owner', 'Due', 'Status']] + [[str(i), '[Action in one line]', '[Name]', '[DD.MM.]', ''] for i in range(1, 8)]
    table(s, L, 1.5, W, rows, colw=[0.6, 6.4, 2.2, 1.6, 1.8], rowh=0.55, size=14)
    status_col(s, 1.5, 0.55, 7, L + W * (10.8 / 12.6) + 0.15, list('ggyyrnn'))

    s = cs(D, C, 'Decision log', 'im_decisions', N('Documented project decisions – prevents re-discussion.', 'Add every decision of the steering committee and project leads.'),
           title='Decision log', source='As of: [date]')
    rows = [['#', 'Date', 'Decision', 'Decided by', 'Impact']] + [['D-%02d' % i, '[DD.MM.]', '[Decision in one sentence]', '[Body / name]', '[Scope / time / cost]'] for i in range(1, 7)]
    table(s, L, 1.5, W, rows, colw=[0.9, 1.4, 5.4, 2.2, 2.7], rowh=0.6, size=14)

    s = cs(D, C, 'Resource plan', 'im_resources', N('Planned capacity per role over time.', 'Edit chart data (FTE per month).'), source='Source: [Resource plan]; FTE')
    bar(s, L, 1.4, 9.0, 4.45, ['M1', 'M2', 'M3', 'M4', 'M5', 'M6', 'M7', 'M8', 'M9', 'M10', 'M11', 'M12'],
        [('Körber', [2, 3, 4, 4, 5, 5, 5, 4, 4, 3, 2, 1]), ('Customer', [1, 2, 2, 3, 3, 4, 4, 4, 3, 3, 2, 2]), ('Vendors', [0, 0, 1, 2, 3, 3, 3, 2, 2, 1, 1, 0])],
        stacked=True, colors=[BLUE, TAUPE, STONE], legend='b', gap=50, size=12, labels=False, val_axis=True)
    kpi(s, 9.9, 1.6, R - 9.9, '00 FTE', 'Peak in month [x]', bigsize=28)
    kpi(s, 9.9, 3.2, R - 9.9, '000 days', 'Customer effort in total', bigsize=28)

    s = cs(D, C, 'Change impact assessment', 'im_change_impact', N('How strongly each stakeholder group is affected – basis for change measures.',
                                                                    'Harvey balls 0–4 per dimension.', 'High impact + low readiness = focus group for change management.'))
    dims = ['Processes', 'Roles', 'Systems', 'Skills']
    groups = ['Operators', 'Shift leads', 'QA', 'Maintenance', 'Management']
    data = [['Stakeholder group'] + dims + ['Measures']] + [[g, '', '', '', '', '[Measure]'] for g in groups]
    table(s, L, 1.5, W, data, colw=[2.6, 1.4, 1.4, 1.4, 1.4, 4.4], rowh=0.65, size=14, aligns=['l', 'c', 'c', 'c', 'c', 'l'])
    xs = [L + W * sum([2.6, 1.4, 1.4, 1.4, 1.4, 4.4][:k]) / 12.6 for k in range(7)]
    sc = [[4, 2, 4, 3], [3, 3, 3, 3], [2, 1, 3, 2], [2, 2, 2, 3], [1, 1, 1, 1]]
    for r in range(5):
        for c in range(4):
            harvey(s, (xs[1 + c] + xs[2 + c]) / 2 - 0.18, 1.5 + (r + 1) * 0.65 + 0.14, 0.36, sc[r][c])

    s = cs(D, C, 'Change approach', 'im_adkar', N('Change measures along the ADKAR model.', 'Adapt measures; assign owners.'),
           title='Change approach along the five ADKAR stages')
    st = [('Awareness', 'Why change?'), ('Desire', 'Why me?'), ('Knowledge', 'How to change?'), ('Ability', 'Can I do it?'), ('Reinforcement', 'Will it stick?')]
    chevrons(s, [a for a, _ in st], L, 1.5, W, 0.65, size=14)
    cw = (W - 4 * 0.06) / 5
    for i, (a, q) in enumerate(st):
        x = L + i * (cw + 0.06)
        text(s, x + 0.1, 2.3, cw - 0.2, 0.35, '„' + q + '“', size=14, italic=True, color=DGRAY)
        rect(s, x, 2.8, cw, 2.9, fill=LSAND)
        text(s, x + 0.15, 2.95, cw - 0.3, 2.6, bl(['Measure', 'Measure', 'Measure'], sa=8), size=14)

    s = cs(D, C, 'Training plan', 'im_training', N('Training per target group before go-live.', 'Add dates and participant numbers; track completion.'), source='As of: [date]')
    rows = [['Target group', 'Training', 'Format', 'Duration', 'Date', 'Participants', 'Completed']] + \
           [[g, '[Training]', f, '[x h]', '[DD.MM.]', '[0]', '00%'] for g, f in
            [('Operators', 'On the job'), ('Shift leads', 'Classroom'), ('QA', 'Workshop'), ('Maintenance', 'Classroom'), ('Key users', 'Train the trainer')]]
    tb = table(s, L, 1.5, W, rows, colw=[2.2, 2.6, 2.0, 1.3, 1.5, 1.5, 1.5], rowh=0.65, size=14, aligns=['l', 'l', 'l', 'c', 'c', 'c', 'c'])
    for r in range(1, 6):
        fill_tf(tb.cell(r, 6).text_frame, '00%', size=14, color=BLUE, align='c')

    s = cs(D, C, 'Go-live readiness', 'im_golive', N('Go / no-go decision before go-live.', 'Status per criterion; the recommendation goes into the title.',
                                                     'Define go-live criteria four weeks before go-live – not in the go/no-go meeting.'),
           title='Go-live readiness: recommendation [go / no-go]')
    crit = ['System qualified and released', 'Data migrated and verified', 'Users trained', 'SOPs effective', 'Support and hypercare organised', 'Rollback plan tested']
    for i, c in enumerate(crit):
        y = 1.5 + i * 0.7
        rect(s, L, y, 7.6, 0.6, fill=LSAND)
        rag(s, L + 0.2, y + 0.19, 'gggygg'[i], 0.22)
        text(s, L + 0.6, y, 5.0, 0.6, c, size=15, anchor='m')
        text(s, L + 5.6, y, 1.9, 0.6, '[Owner]', size=13, color=DGRAY, anchor='m')
    xr = 8.4; wr = R - xr
    rect(s, xr, 1.5, wr, 4.1, fill=BLUE)
    label(s, xr + 0.3, 1.75, wr - 0.6, 'Recommendation', color=WHITE)
    text(s, xr + 0.3, 2.1, wr - 0.6, 1.5, 'Go-live on [date] with [condition]', size=22, color=WHITE)
    rrect(s, xr + 0.3, 4.75, 2.6, 0.5, fill=YELLOW, radius=0.25, content='Decision: go / no-go', size=13, align='c', anchor='m', margin=0)

    s = cs(D, C, 'Cut-over plan', 'im_cutover', N('Detailed plan for the switch from old to new.', 'Adapt time slots (hours or days) and activities.'), source='As of: [date]')
    slots = ['T-3', 'T-2', 'T-1', 'Go-live', 'T+1', 'T+2', 'T+3']
    rows = [('Production stop / freeze', 0, 1, 'bar'), ('Data migration', 1, 2.5, 'hl'), ('Verification and release', 2.5, 3, 'hl'), ('Go-live', 3, 0, 'ms'),
            ('First batch', 3, 4, 'hl'), ('Hypercare', 3, 7, 'bar'), ('Rollback window', 3, 5, 'bar')]
    gantt(s, L, 1.5, W, rows, slots, lab_w=3.4, rowh=0.55, size=14)


# =====================================================================================
def workshops(D):
    C = 'Workshops'
    divider(D, 0, 'Workshops', N('Facilitation slides for customer workshops.', 'Combine with the agenda slides from chapter 01.',
                                 'A workshop without a defined output is a meeting – write the expected result on the first slide.'))
    s = cs(D, C, 'Objectives and agenda', 'ws_objectives', N('First slide of every workshop.', 'Define output and agenda; plan 20% buffer.'),
           title='Workshop [topic]: objectives and agenda')
    rect(s, L, 1.5, 5.4, 4.3, fill=BLUE)
    label(s, L + 0.3, 1.75, 5, 'Today’s expected output', color=WHITE)
    text(s, L + 0.3, 2.1, 4.8, 3.4, bl(['Output 1', 'Output 2', 'Output 3'], sa=10), size=18, color=WHITE)
    xr = 6.3; wr = R - xr
    for i in range(6):
        y = 1.5 + i * 0.72
        text(s, xr, y, 1.5, 0.6, '00:00', size=15, color=BLUE, anchor='m')
        text(s, xr + 1.6, y, wr - 1.6, 0.6, 'Session %d' % (i + 1) if i != 3 else 'Break', size=16, anchor='m', color=DGRAY if i == 3 else BLACK)
        hline(s, xr, y + 0.68, wr, STONE, 0.5)

    s = cs(D, C, 'Exercise / breakout', 'ws_exercise', N('Instruction for a group exercise.', 'Task, groups, time and expected output.'),
           title='Exercise: [task in one line]')
    disturber(s, L, 1.6, [{'t': '00', 'size': 48}, {'t': 'minutes', 'size': 14}], d=2.4)
    x0 = L + 3.0; cw = (R - x0 - 2 * 0.3) / 3
    for i, (h, ic) in enumerate([('Task', 'assignment'), ('Groups', 'groups'), ('Output', 'flag')]):
        colhead(s, x0 + i * (cw + 0.3), 1.6, cw, None, h, bl(['Point', 'Point', 'Point'], sa=6), icon_=ic, hsize=19, bsize=15, bh=2.0)

    s = cs(D, C, 'Brainstorming board', 'ws_board', N('Clustered ideas from a brainstorming session.', 'Copy sticky notes; one idea per note.'))
    cw = (W - 3 * 0.3) / 4
    for i in range(4):
        x = L + i * (cw + 0.3)
        label(s, x, 1.45, cw, 'Cluster %d' % (i + 1))
        hline(s, x, 1.75, cw, BLACK)
        for k in range(3 if i % 2 else 4):
            sx = x + (k % 2) * (cw / 2 + 0.05); sy = 1.9 + (k // 2) * 1.35
            sp = rect(s, sx, sy, cw / 2 - 0.05, 1.2, fill=LBLUE if i % 2 == 0 else SAND, content='Idea', size=13, anchor='m', align='c')
            sp.rotation = [-2, 1.5, -1, 2][k]

    s = cs(D, C, 'Voting results', 'ws_voting', N('Result of dot voting or prioritisation.', 'Edit chart data; sort descending; top 3 in blue.'),
           source='Source: Workshop [date], n = [xx] participants')
    bar(s, L, 1.4, 8.6, 4.45, ['Idea %d' % i for i in range(1, 9)], [('Votes', [14, 11, 9, 6, 5, 3, 2, 1])], kind='bar', fmt='0', highlight=[0, 1, 2], gap=35, size=13)
    label(s, 9.5, 1.5, R - 9.5, 'Next steps for the top 3')
    text(s, 9.5, 1.85, R - 9.5, 3, bl(['Idea 1: [owner, date]', 'Idea 2: [owner, date]', 'Idea 3: [owner, date]'], sa=8), size=14)

    s = cs(D, C, 'Parking lot', 'ws_parking', N('Topics and questions that came up but are not part of the session.', 'Assign an owner before the workshop ends.'),
           title='Parking lot and open questions')
    cw = (W - 0.5) / 2
    for i, (h, ic) in enumerate([('Parking lot', 'local_shipping' if False else 'inventory_2'), ('Open questions', 'help')]):
        x = L + i * (cw + 0.5)
        icon(s, ic, x, 1.5, 0.42, BLUE)
        text(s, x + 0.6, 1.52, cw - 0.6, 0.4, h, size=19)
        hline(s, x, 2.1, cw, BLACK)
        for k in range(5):
            y = 2.2 + k * 0.68
            text(s, x, y, cw - 1.8, 0.58, 'Topic / question', size=15, anchor='m')
            text(s, x + cw - 1.8, y, 1.8, 0.58, '[Owner]', size=13, color=DGRAY, anchor='m', align='r')
            hline(s, x, y + 0.63, cw, STONE, 0.5)

    s = D.slide('hl_dark', 'Break', key='ws_break', notes=N('Break slide for workshops.', 'Change the return time in the text.'))
    pic(s, 27, IMG + 'image3.png')
    D.ph(s, 29, 'until 00:00')
    D.drop_empty(s, keep=(20, 21, 22, 23, 24, 27, 29))


# =====================================================================================
def appendix(D):
    C = 'Appendix'
    divider(D, 0, 'Appendix', N('Back-up, glossary and methodology.', 'Use after the closing slide; reference back-up slides in the main part.',
                                'Move detail into the appendix – the main storyline stays short.'))
    s = cs(D, C, 'Back-up data table', 'ap_backup', N('Detailed data table for the appendix.', 'Up to 10 rows; reference it from the main slide.'), source='Source: [Source]',
           sticker_='Back-up')
    rows = [['Item', 'Value 1', 'Value 2', 'Value 3', 'Value 4', 'Comment']] + [['Item %d' % i, '0', '0', '0', '0', '[Comment]'] for i in range(1, 9)]
    table(s, L, 1.5, W, rows, colw=[3.0, 1.5, 1.5, 1.5, 1.5, 3.6], rowh=0.48, size=13, aligns=['l', 'r', 'r', 'r', 'r', 'l'])

    s = cs(D, C, 'Glossary', 'ap_glossary', N('Abbreviations used in pharma projects.', 'Delete unused terms; add customer-specific ones.'),
           title='Glossary of common abbreviations')
    terms = [('CAPA', 'Corrective and preventive action'), ('CCS', 'Contamination control strategy'), ('CQV', 'Commissioning, qualification, validation'),
             ('CSV / CSA', 'Computer system validation / assurance'), ('eBR', 'Electronic batch record'), ('FAT / SAT', 'Factory / site acceptance test'),
             ('GAMP', 'Good automated manufacturing practice'), ('GxP', 'Good practice regulations (GMP, GDP …)'), ('IQ / OQ / PQ', 'Installation / operational / performance qualification'),
             ('MES', 'Manufacturing execution system'), ('OEE', 'Overall equipment effectiveness'), ('URS', 'User requirement specification')]
    cw = (W - 0.5) / 2
    for i, (a, b) in enumerate(terms):
        x = L + (i // 6) * (cw + 0.5); y = 1.5 + (i % 6) * 0.68
        text(s, x, y, 1.9, 0.6, a, size=15, color=BLUE, anchor='m')
        text(s, x + 2.0, y, cw - 2.0, 0.6, b, size=14, anchor='m')
        hline(s, x, y + 0.64, cw, STONE, 0.5)

    s = cs(D, C, 'Methodology and sources', 'ap_method', N('How the analysis was done – builds credibility.', 'List data sources, sample sizes and periods.'))
    rows = [['Analysis', 'Method', 'Data source', 'Period', 'Sample']] + [['Analysis %d' % i, '[Method]', '[Source]', '[MM/YY – MM/YY]', 'n = [xx]'] for i in range(1, 6)]
    table(s, L, 1.5, W, rows, colw=[2.6, 2.8, 3.0, 2.4, 1.8], rowh=0.65, size=14)
