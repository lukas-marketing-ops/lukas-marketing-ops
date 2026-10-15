from en import *
from pptx.enum.dml import MSO_LINE


# =====================================================================================
def s05(D):
    C = 'Solution & value'
    divider(D, 5, 'Solution\n& value', N('Value proposition, solution, options and recommendation.', 'Always argue from the customer benefit.',
                                         'Logic: problem → solution → measurable benefit → recommendation.'))
    s = cs(D, C, 'Value proposition', 'valueprop', N('Core message of the solution presentation.', 'Promise in one sentence, three proof points.',
                                                     'The statement must work without the slide – as an elevator pitch.'))
    rect(s, L, 1.5, 5.6, 4.3, fill=BLUE)
    label(s, L + 0.35, 1.8, 5, 'Our promise', color=WHITE)
    text(s, L + 0.35, 2.15, 4.9, 3.0, 'Value proposition in one sentence: what the customer achieves with us.', size=24, color=WHITE)
    xr = 6.5; wr = R - xr
    for i in range(3):
        y = 1.5 + i * 1.45
        hline(s, xr, y, wr, BLACK)
        text(s, xr, y + 0.15, 2.4, 0.7, '00%', size=32, color=BLUE)
        text(s, xr + 2.6, y + 0.25, wr - 2.6, 0.9, 'Proof point that supports the promise', size=15)

    s = cs(D, C, 'Solution modules', 'modules', N('Modular offer – makes it easy to start small.', 'Adapt duration and deliverables; recommended module in blue.',
                                                  'Land and expand: start with module 1, deliver results, then extend.'))
    cw = (W - 3 * 0.3) / 4
    for i in range(4):
        x = L + i * (cw + 0.3)
        f = BLUE if i == 1 else LSAND; c = WHITE if i == 1 else BLACK
        rect(s, x, 1.5, cw, 4.25, fill=f)
        label(s, x + 0.25, 1.75, cw - 0.5, 'Module %d' % (i + 1), color=WHITE if i == 1 else BLUE)
        text(s, x + 0.25, 2.05, cw - 0.5, 0.45, 'Module name', size=21, color=c)
        text(s, x + 0.25, 2.6, cw - 0.5, 0.3, '[Duration]', size=13, color=WHITE if i == 1 else DGRAY)
        hline(s, x + 0.25, 3.05, cw - 0.5, WHITE if i == 1 else STONE, 0.75)
        text(s, x + 0.25, 3.2, cw - 0.5, 1.6, bl(['Deliverable', 'Deliverable', 'Deliverable'], sa=6), size=14, color=c)
        if i < 3:
            icon(s, 'arrow_forward', x + cw + 0.04, 3.45, 0.22, TAUPE)

    s = cs(D, C, 'Value driver tree', 'valuetree', N('Translates operational levers into financial impact – for CFO and management.',
                                                      'Every leaf needs a KPI and an owner at the customer.'), source='Source: [Source]')
    shape(s, SH.RECTANGLE, L, 3.05, 2.5, 1.2, fill=BLUE, content=[{'t': 'Total impact p.a.', 'size': 13}, {'t': '€0.0m', 'size': 24}], color=WHITE, anchor='m')
    for i, d in enumerate(['Revenue', 'Cost', 'Capital']):
        y = 1.55 + i * 1.65
        arrow(s, L + 2.5, 3.65, L + 3.1, y + 0.45, color=STONE)
        shape(s, SH.RECTANGLE, L + 3.1, y + 0.1, 2.5, 0.75, fill=SAND, content=[{'t': d, 'size': 14}, {'t': '€0.0m', 'size': 15, 'color': BLUE}], anchor='m')
        for j in range(2):
            yy = y + j * 0.78 - 0.15
            arrow(s, L + 5.6, y + 0.47, L + 6.2, yy + 0.32, color=STONE)
            shape(s, SH.RECTANGLE, L + 6.2, yy, 3.6, 0.64, line=STONE, content='Lever', size=14, anchor='m')
            text(s, L + 10.0, yy, 1.3, 0.64, '[KPI]', size=13, anchor='m', color=DGRAY)
            text(s, L + 11.3, yy, 1.3, 0.64, '€0.0m', size=15, anchor='m', color=BLUE, align='r')

    s = cs(D, C, 'Benefits and tracking', 'kpis', N('Links the promise to a measurement logic.', 'Agree KPI, baseline and method with the customer.',
                                                    'No baseline, no proof: measure it before the project starts.'))
    cw = (W - 3 * 0.3) / 4
    for i in range(4):
        x = L + i * (cw + 0.3)
        rect(s, x, 1.5, cw, 2.5, fill=BLUE if i == 3 else LSAND)
        col = WHITE if i == 3 else BLUE
        text(s, x + 0.25, 1.75, cw - 0.5, 0.8, ['+00%', '–00%', '–0 mo.', '€0.0m'][i], size=36, color=col)
        text(s, x + 0.25, 2.7, cw - 0.5, 0.35, 'KPI name', size=16, color=WHITE if i == 3 else BLACK)
        text(s, x + 0.25, 3.1, cw - 0.5, 0.35, 'Baseline → target', size=13, color=WHITE if i == 3 else DGRAY)
        label(s, x, 4.3, cw, 'Measured by', color=DGRAY)
        hline(s, x, 4.6, cw, STONE, 0.5)
        text(s, x, 4.72, cw, 0.6, '[Source, frequency]', size=13)
    sowhat(s, KT)

    s = cs(D, C, 'Today and target picture', 'asis_tobe', N('Before-and-after comparison.', 'Maximum five rows.'))
    label(s, L + 2.4, 1.45, 4.4, 'Today', color=DGRAY); label(s, L + 7.4, 1.45, 5, 'Target picture', color=BLUE)
    for i in range(4):
        y = 1.8 + i * 0.95
        text(s, L, y, 2.3, 0.8, 'Dimension %d' % (i + 1), size=16, anchor='m')
        rect(s, L + 2.4, y, 4.4, 0.8, fill=LSAND, content='Current situation', size=14, anchor='m', margin=(0.2, 0.03, 0.1, 0.03))
        icon(s, 'arrow_forward', L + 6.95, y + 0.27, 0.26, BLUE)
        rect(s, L + 7.4, y, R - L - 7.4, 0.8, fill=BLUE, content='Target situation', size=14, color=WHITE, anchor='m', margin=(0.2, 0.03, 0.1, 0.03))
    sowhat(s, KT)

    s = cs(D, C, 'Options', 'options', N('Option comparison before the recommendation.', 'Recommended option in blue; always show at least three.'))
    cw = (W - 2 * 0.3) / 3
    for i in range(3):
        x = L + i * (cw + 0.3)
        hi = i == 1; f = BLUE if hi else LSAND; c = WHITE if hi else BLACK
        rect(s, x, 1.5, cw, 4.25, fill=f)
        label(s, x + 0.3, 1.75, cw - 0.6, 'Recommended' if hi else 'Option', color=WHITE if hi else BLUE)
        text(s, x + 0.3, 2.05, cw - 0.6, 0.45, 'Option %s' % 'ABC'[i], size=24, color=c)
        text(s, x + 0.3, 2.6, cw - 0.6, 0.4, D1, size=15, color=c)
        hline(s, x + 0.3, 3.35, cw - 0.6, WHITE if hi else STONE, 0.75)
        text(s, x + 0.3, 3.5, cw - 0.6, 1.4, bl(['Investment', 'Benefit', 'Timing'], sa=6), size=15, color=c)

    s = cs(D, C, 'Evaluation matrix', 'evaluation', N('Weighted evaluation of options.', 'Agree criteria and weights with the customer first; Harvey balls 0–4.'))
    x0 = L + 4.2; cw = 2.6
    for j in range(3):
        if j == 1:
            rect(s, x0 + j * cw, 1.4, cw, 4.4, fill=LBLUE)
        text(s, x0 + j * cw, 1.5, cw, 0.3, 'Option %s' % 'ABC'[j], size=15, color=BLUE, align='c')
    text(s, L + 2.9, 1.5, 1.2, 0.3, 'Weight', size=13, color=DGRAY, align='c')
    hline(s, L, 1.9, W, BLACK, 1)
    sc = [(2, 4, 2), (4, 2, 3), (3, 3, 2), (2, 4, 1), (1, 3, 3)]
    for i, row in enumerate(sc):
        y = 2.0 + i * 0.62
        text(s, L, y, 2.8, 0.55, 'Criterion %d' % (i + 1), size=15, anchor='m')
        text(s, L + 2.9, y, 1.2, 0.55, '00%', size=14, anchor='m', align='c', color=DGRAY)
        for j, q in enumerate(row):
            harvey(s, x0 + j * cw + cw / 2 - 0.18, y + 0.09, 0.36, q)
        hline(s, L, y + 0.6, W, STONE, 0.5)
    y = 2.0 + 5 * 0.62 + 0.1
    text(s, L, y, 2.8, 0.5, 'Weighted score', size=15, anchor='m')
    for j in range(3):
        text(s, x0 + j * cw, y, cw, 0.5, '0.0', size=20, align='c', anchor='m', color=BLUE if j == 1 else BLACK)

    s = cs(D, C, 'Recommendation', 'recommendation', N('Clear recommendation with a call to action.', 'Name decision, budget and date.',
                                                       'Put the ask into the title – decision makers must know what is expected.'))
    rect(s, L, 1.5, 5.6, 4.3, fill=BLUE)
    label(s, L + 0.35, 1.8, 5, 'Our recommendation', color=WHITE)
    text(s, L + 0.35, 2.15, 4.9, 2.2, 'Recommendation and decision required in up to three lines.', size=24, color=WHITE)
    rrect(s, L + 0.35, 4.95, 3.3, 0.48, fill=YELLOW, radius=0.24, content='Decision by [date]', size=14, color=BLACK, align='c', anchor='m', margin=0)
    xr = 6.5; wr = R - xr
    label(s, xr, 1.5, wr, 'Why we recommend this')
    for i in range(3):
        y = 1.9 + i * 1.25
        icon(s, 'check_circle', xr, y, 0.42, BLUE)
        text(s, xr + 0.6, y, wr - 0.6, 0.35, 'Reason %d' % (i + 1), size=17)
        text(s, xr + 0.6, y + 0.42, wr - 0.6, 0.6, D1, size=14, color=DGRAY)


# =====================================================================================
def s06(D):
    C = 'Business case'
    divider(D, 6, 'Business\ncase', N('Investment logic: summary, cash flow, value bridge, scenarios.', 'Validate figures with the customer’s controlling.',
                                       'Build the business case together with the customer – then they defend it internally.'))
    s = cs(D, C, 'Summary', 'bc_summary', N('One-pager for the investment decision.', 'Copy values from the business-case model.'), source='Source: [Model, version]')
    cw = (W - 0.5) / 2
    for side, lab in enumerate(['Investment (one-off)', 'Benefits (p.a.)']):
        x = L + side * (cw + 0.5)
        label(s, x, 1.5, cw, lab, color=BLUE)
        hline(s, x, 1.82, cw, BLACK)
        for i in range(3):
            y = 1.9 + i * 0.45
            text(s, x, y, cw - 1.6, 0.4, 'Item %d' % (i + 1), size=14, anchor='m'); text(s, x + cw - 1.6, y, 1.6, 0.4, '€0.0m', size=14, anchor='m', align='r')
            hline(s, x, y + 0.43, cw, STONE, 0.5)
        y = 1.9 + 3 * 0.45 + 0.05
        text(s, x, y, cw - 1.6, 0.4, 'Total', size=15, anchor='m'); text(s, x + cw - 1.6, y, 1.6, 0.4, '€0.0m', size=15, anchor='m', align='r', color=BLUE)
    cw3 = (W - 2 * 0.3) / 3
    for i, (a, b) in enumerate([('€0.0m', 'Net present value (NPV)'), ('0.0 years', 'Payback period'), ('00%', 'Internal rate of return (IRR)')]):
        x = L + i * (cw3 + 0.3)
        rect(s, x, 3.95, cw3, 1.85, fill=BLUE if i == 0 else LSAND)
        text(s, x + 0.3, 4.15, cw3 - 0.6, 0.8, a, size=36, color=WHITE if i == 0 else BLUE)
        text(s, x + 0.3, 5.05, cw3 - 0.6, 0.4, b, size=14, color=WHITE if i == 0 else BLACK)

    s = cs(D, C, 'Cash flow', 'bc_cashflow', N('Investment phase and break-even.', 'Edit data; negative values grey, positive blue.'), source='Source: [Model]; €m')
    vals = [-6, -14, -10, -4, 2, 9, 15]
    ch = bar(s, L, 1.4, 8.4, 4.4, ['Year 1', 'Year 2', 'Year 3', 'Year 4', 'Year 5', 'Year 6', 'Year 7'], [('Cumulative cash flow', vals)],
             fmt='0', point_colors=[STONE if v < 0 else BLUE for v in vals], gap=50, size=13)
    ch.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    xr = 9.3; wr = R - xr
    kpi(s, xr, 1.6, wr, '€0.0m', 'Maximum cash need', bigsize=30)
    kpi(s, xr, 3.2, wr, 'Year 0', 'Break-even', bigsize=30)
    sowhat(s, KT)

    s = cs(D, C, 'Value bridge', 'bc_waterfall', N('Waterfall: where the value comes from.', 'Edit data: columns “Summe” (totals), “Zunahme” (up), “Abnahme” (down); “Basis” stays invisible.'),
           source='Source: [Model]; present values in €m')
    waterfall(s, L, 1.4, W, 4.4, ['Lever 1', 'Lever 2', 'Lever 3', 'Gross benefit', 'Operating cost', 'Investment', 'NPV'],
              [12.0, 6.0, 4.0, 22.0, -3.0, -14.0, 5.0], totals=[3, 6], fmt='{:+.0f}', tfmt='{:.0f}', size=13, colors=[None, BLUE, STONE, BLACK])
    sowhat(s, KT)

    s = cs(D, C, 'Scenarios', 'scenarios', N('Scenario analysis for the investment decision.', 'Document the assumptions per scenario.'), source='Source: [Model]')
    rows = [['Assumption', 'Conservative', 'Base', 'Optimistic']] + [['Assumption %d' % i, '[Value]', '[Value]', '[Value]'] for i in range(1, 4)] + \
           [[('NPV', {}), ('€0.0m', {'size': 17}), ('€0.0m', {'size': 17, 'color': WHITE, 'fill': BLUE}), ('€0.0m', {'size': 17})]]
    table(s, L, 1.5, W, rows, colw=[4, 2.8, 2.8, 2.8], rowh=0.65, size=15, aligns=['l', 'c', 'c', 'c'], fills={(r, 2): LBLUE for r in range(1, 4)})
    sowhat(s, KT)


# =====================================================================================
def s07(D):
    C = 'Approach & plan'
    divider(D, 7, 'Approach\n& plan', N('Project approach, timeline, organisation and governance.', 'Adapt phases, dates and roles to the project.',
                                          'Concrete beats generic: the approach shows that we master the project.'))
    s = cs(D, C, 'Phase model', 'phases', N('Overview of the approach.', 'Phrase deliverables as nouns – they are the basis for payment milestones.'))
    chevrons(s, ['Phase 1', 'Phase 2', 'Phase 3', 'Phase 4'], L + 1.8, 1.5, W - 1.8, 0.62, size=14)
    cw = (W - 1.8 - 3 * 0.06) / 4
    y = 2.3
    for j, (lab, v, h_) in enumerate([('Duration', '[x weeks]', 0.55), ('Focus', 'Key question of the phase', 1.0), ('Deliverable', 'Deliverable', 0.9)]):
        hline(s, L, y, W, STONE, 0.5)
        label(s, L, y + 0.15, 1.7, lab, color=BLUE if j == 2 else DGRAY)
        for i in range(4):
            x = L + 1.8 + i * (cw + 0.06)
            if j == 2:
                rect(s, x, y + 0.08, cw, h_ - 0.16, fill=LSAND)
            text(s, x + 0.15, y + 0.08, cw - 0.3, h_ - 0.16, v, size=15, anchor='m')
        y += h_
    sowhat(s, KT)

    s = cs(D, C, 'Phase detail', 'phasedetail', N('Detail of one phase.', 'Duplicate per phase; mark the phase in the chevron row.'))
    chevrons(s, ['Phase 1', 'Phase 2', 'Phase 3', 'Phase 4'], L, 1.45, W, 0.45, active=0, size=12)
    cw = (W - 3 * 0.35) / 4
    for i, (h, ic) in enumerate([('Objectives', 'target'), ('Activities', 'checklist'), ('Deliverables', 'description'), ('Your involvement', 'groups')]):
        colhead(s, L + i * (cw + 0.35), 2.2, cw, None, h, bl(B3, sa=6), icon_=ic, hsize=19, bsize=14, bh=1.8)
    rect(s, L, 5.2, W, 0.6, fill=LSAND)
    text(s, L + 0.25, 5.2, W - 0.5, 0.6, [{'t': [('Duration  ', {'color': BLUE}), ('[x weeks]      ', {}), ('Team  ', {'color': BLUE}), ('[roles]      ', {}),
                                                    ('Gate  ', {'color': BLUE}), ('[decision forum, date]', {})]}], size=14, anchor='m')

    s = cs(D, C, 'Gantt chart', 'gantt', N('Timeline on month level.', 'Move bars; critical path in blue. Maximum 10 rows.'), source='As of: [date]')
    months = ['J', 'F', 'M', 'A', 'M', 'J', 'J', 'A', 'S', 'O', 'N', 'D']
    text(s, L + 3.0, 1.38, 9.6, 0.25, '[Year]', size=11, color=BLUE)
    rows = [('Workstream 1', 0, 2, 'hl'), ('Workstream 2', 1.5, 4, 'hl'), ('Milestone', 4, 0, 'ms'), ('Workstream 3', 4, 7, 'hl'),
            ('Workstream 4', 5, 8, 'bar'), ('Workstream 5', 7.5, 10, 'hl'), ('Milestone', 10, 0, 'ms'), ('Workstream 6', 9, 12, 'bar')]
    gantt(s, L, 1.6, W, rows, months, lab_w=3.0, rowh=0.48, size=14)
    rect(s, L, 6.25, 0.3, 0.16, fill=BLUE); text(s, L + 0.4, 6.2, 2, 0.25, 'Critical path', size=11)
    rect(s, L + 2.0, 6.25, 0.3, 0.16, fill=STONE); text(s, L + 2.4, 6.2, 2, 0.25, 'Other activities', size=11)
    shape(s, SH.DIAMOND, L + 4.3, 6.21, 0.22, 0.22, fill=BLACK); text(s, L + 4.62, 6.2, 2, 0.25, 'Milestone', size=11)

    s = cs(D, C, 'Milestones', 'milestones', N('Milestone timeline for management.', 'Filled = reached, outline = planned.'))
    y0 = 3.55; cw = W / 5
    hline(s, L, y0, W, BLACK, 1.25)
    for i in range(5):
        x = L + i * cw
        oval(s, x + 0.02, y0 - 0.15, 0.3, fill=BLUE if i < 2 else WHITE, line=BLUE, lw=1.75)
        yy = 1.75 if i % 2 == 0 else 3.95
        text(s, x, yy, cw - 0.2, 0.5, '[Q0 YYYY]', size=24, color=BLUE)
        text(s, x, yy + 0.6, cw - 0.2, 0.35, 'Milestone', size=16)
        text(s, x, yy + 0.98, cw - 0.2, 0.5, D1, size=13, color=DGRAY)
    dot(s, L, 6.25, 0.18, BLUE); text(s, L + 0.28, 6.2, 1.5, 0.3, 'Reached', size=11)
    oval(s, L + 1.4, 6.25, 0.18, fill=WHITE, line=BLUE, lw=1.5); text(s, L + 1.68, 6.2, 1.5, 0.3, 'Planned', size=11)

    s = cs(D, C, 'Stage gates', 'stagegate', N('Shows risk control through decision gates – convinces CFO and procurement.', 'Define gate criteria per project.'))
    pw = (W - 4 * 0.55) / 4
    for i in range(4):
        x = L + i * (pw + 0.55)
        rect(s, x, 1.6, pw, 0.62, fill=BLUE if i == 0 else SAND, content='Phase %d' % (i + 1), size=15, color=WHITE if i == 0 else BLACK, align='c', anchor='m')
        shape(s, SH.DIAMOND, x + pw + 0.1, 1.6, 0.36, 0.62, fill=BLACK)
        label(s, x, 2.6, pw + 0.5, 'Gate %d – criteria' % (i + 1), color=BLUE)
        hline(s, x, 2.9, pw + 0.45, STONE, 0.5)
        text(s, x, 3.0, pw + 0.4, 1.2, bl(['Criterion', 'Criterion'], sa=6), size=14)
    rect(s, L, 4.6, W, 1.2, fill=LSAND)
    text(s, L + 0.3, 4.6, W - 0.6, 1.2, [{'t': [('Go  ', {'color': BLUE}), ('start next phase      ', {}), ('Hold  ', {'color': BLUE}),
                                                ('rework with new date      ', {}), ('Stop  ', {'color': BLUE}), ('end project, secure results', {})]}], size=15, anchor='m')

    s = cs(D, C, 'Workshop concept', 'workshops', N('Shows when and how the customer is involved.', 'Adapt workshops, participants and duration.'))
    cw = (W - 3 * 0.3) / 4
    for i in range(4):
        x = L + i * (cw + 0.3)
        oval(s, x, 1.5, 0.8, fill=BLUE, content='W%d' % (i + 1), size=17, color=WHITE, align='c', anchor='m', margin=0)
        text(s, x + 0.95, 1.55, cw - 0.95, 0.35, 'Workshop', size=18)
        text(s, x + 0.95, 1.95, cw - 0.95, 0.3, '[Week]', size=13, color=DGRAY)
        hline(s, x, 2.55, cw, BLACK)
        for k, (lab, v) in enumerate([('Objective', D1), ('Participants', '[Roles]'), ('Duration', '[x hours]')]):
            label(s, x, 2.7 + k * 1.0, cw, lab); text(s, x, 3.0 + k * 1.0, cw, 0.65, v, size=14)
        if i < 3:
            icon(s, 'arrow_forward', x + cw + 0.03, 1.77, 0.24, TAUPE)

    s = cs(D, C, 'Project organisation', 'org', N('Project organisation for proposal and kick-off.', 'Fill names; tandem leads on both sides.'))
    rect(s, 3.67, 1.45, 6.0, 0.75, fill=BLUE, content=[{'t': 'Steering committee', 'size': 16}, {'t': '[Names]', 'size': 12}], color=WHITE, align='c', anchor='m')
    vline(s, 6.67, 2.2, 0.35, STONE)
    rect(s, 4.17, 2.55, 2.4, 0.8, fill=SAND, content=[{'t': 'Project lead customer', 'size': 13}, {'t': '[Name]', 'size': 11}], align='c', anchor='m')
    rect(s, 6.77, 2.55, 2.4, 0.8, fill=SAND, content=[{'t': 'Project lead Körber', 'size': 13}, {'t': '[Name]', 'size': 11}], align='c', anchor='m')
    rect(s, 10.0, 2.55, 2.97, 0.8, line=STONE, content=[{'t': 'PMO', 'size': 13}, {'t': '[Name]', 'size': 11}], align='c', anchor='m')
    hline(s, 9.17, 2.95, 0.83, STONE); vline(s, 6.67, 3.35, 0.4, STONE)
    cw = (W - 3 * 0.25) / 4
    hline(s, L + cw / 2, 3.75, W - cw, STONE)
    for i in range(4):
        x = L + i * (cw + 0.25)
        vline(s, x + cw / 2, 3.75, 0.3, STONE)
        rect(s, x, 4.05, cw, 1.7, line=STONE)
        label(s, x + 0.2, 4.2, cw - 0.4, 'Workstream %d' % (i + 1))
        text(s, x + 0.2, 4.5, cw - 0.4, 0.35, 'Workstream name', size=16)
        text(s, x + 0.2, 4.95, cw - 0.4, 0.7, 'Lead customer: [Name]\nLead Körber: [Name]', size=12, color=DGRAY)

    s = cs(D, C, 'RACI matrix', 'raci', N('Responsibility matrix for kick-off and project handbook.', 'Exactly one “A” per row.'))
    roles = ['Sponsor', 'PL customer', 'PL Körber', 'Role 4', 'Role 5']
    pat = ['A R C I I', 'I A R C C', 'I C R A C', 'A C C R I']
    data = [['Deliverable'] + roles] + [['Deliverable %d' % (i + 1)] + p.split() for i, p in enumerate(pat)]
    tb = table(s, L, 1.5, W, data, colw=[4.1] + [1.7] * 5, rowh=0.6, size=15, aligns=['l'] + ['c'] * 5)
    for r in range(1, 5):
        for c in range(1, 6):
            v = data[r][c]
            if v in ('A', 'R'):
                cell = tb.cell(r, c); cell.fill.solid(); setc(cell.fill.fore_color, BLUE if v == 'A' else LBLUE)
                fill_tf(cell.text_frame, v, size=15, color=WHITE if v == 'A' else BLUE, align='c')
    text(s, L, 4.75, W, 0.3, [{'t': [('R ', {'color': BLUE}), ('Responsible    ', {}), ('A ', {'color': BLUE}), ('Accountable    ', {}),
                                      ('C ', {'color': BLUE}), ('Consulted    ', {}), ('I ', {'color': BLUE}), ('Informed', {})]}], size=12)

    s = cs(D, C, 'Governance', 'governance', N('Steering and communication structure.', 'Adapt bodies, participants and rhythm.'))
    rows = [['Body', 'Participants', 'Rhythm', 'Purpose']] + [[b, '[Roles]', r, '[Purpose]'] for b, r in
                                                            [('Steering committee', 'Monthly'), ('Project leads', 'Weekly'), ('Workstream meeting', 'Weekly')]]
    table(s, L, 1.5, W, rows, colw=[3.2, 3.4, 2.2, 3.8], rowh=0.75, size=15)
    sowhat(s, KT)


# =====================================================================================
def s08(D):
    C = 'Processes'
    divider(D, 8, 'Processes', N('Standard process visuals: map, swimlane, SIPOC, process step, value stream, to-be process, decision logic, KPIs.',
                                  'Always use the same shapes: step = rectangle, decision = diamond, waiting time = triangle.',
                                  'A uniform process notation makes documentation comparable across projects.'))
    s = cs(D, C, 'Process map', 'processmap', N('Entry into every process analysis.', 'Mark the processes in scope in blue.'))
    label(s, L, 1.45, 4, 'Management processes', color=DGRAY)
    mw = (W - 2 * 0.2) / 3
    for i in range(3):
        rect(s, L + i * (mw + 0.2), 1.75, mw, 0.65, fill=SAND, content='Process', size=14, align='c', anchor='m')
    label(s, L, 2.7, 4, 'Core processes', color=BLUE)
    chevrons(s, ['Process', 'Process', 'Process', 'Process', 'Process'], L, 3.0, W, 1.1, active=[2], size=15)
    label(s, L, 4.4, 4, 'Support processes', color=DGRAY)
    sw = (W - 3 * 0.2) / 4
    for i in range(4):
        rect(s, L + i * (sw + 0.2), 4.7, sw, 0.65, fill=SAND, content='Process', size=14, align='c', anchor='m')
    rect(s, L, 5.6, 0.22, 0.22, fill=BLUE); text(s, L + 0.32, 5.56, 4, 0.3, 'In scope', size=12)

    s = cs(D, C, 'Swimlane', 'swimlane', N('Process across departments – makes hand-overs visible.', 'Lanes = organisational units; bottlenecks in blue.',
                                           'Count the hand-overs: every lane change is a risk for waiting time.'))
    lh = 1.0; lx = L + 1.7; y0 = 1.45
    for i in range(4):
        y = y0 + i * lh
        rect(s, L, y, 1.6, lh - 0.06, fill=SAND, content='Unit %d' % (i + 1), size=14, anchor='m')
        rect(s, lx, y, R - lx, lh - 0.06, fill=LSAND)
    steps = [(0, 0, False), (1, 1, False), (2, 2, True), (1, 3, False), (2, 4, True), (3, 5, False)]
    sw = 1.45; gap = (R - lx - 0.2 - 6 * sw) / 5 + sw
    pos = {}
    for k, (lane, col, hi) in enumerate(steps):
        x = lx + 0.1 + col * gap; y = y0 + lane * lh + 0.17
        rect(s, x, y, sw, lh - 0.4, fill=BLUE if hi else WHITE, line=None if hi else STONE, content='Step %d' % (k + 1), size=13,
             color=WHITE if hi else BLACK, align='c', anchor='m', margin=0.04)
        pos[k] = (x, y)
    for a in range(5):
        xa, ya = pos[a]; xb, yb = pos[a + 1]
        elbow(s, xa + sw, ya + (lh - 0.4) / 2, xb, yb + (lh - 0.4) / 2)
    sowhat(s, KT)

    s = cs(D, C, 'SIPOC', 'sipoc', N('Process scoping at project start.', '3–5 entries per column.'))
    cw = (W - 4 * 0.2) / 5
    for i, (l_, h) in enumerate([('S', 'Supplier'), ('I', 'Input'), ('P', 'Process'), ('O', 'Output'), ('C', 'Customer')]):
        x = L + i * (cw + 0.2); hi = i == 2
        rect(s, x, 1.5, cw, 0.95, fill=BLUE if hi else SAND)
        text(s, x + 0.2, 1.55, 0.8, 0.85, l_, size=36, color=WHITE if hi else BLUE, anchor='m')
        text(s, x + 0.9, 1.55, cw - 1.0, 0.85, h, size=16, color=WHITE if hi else BLACK, anchor='m')
        for j in range(4 if hi else 3):
            rect(s, x, 2.65 + j * 0.68, cw, 0.56, line=BLUE if hi else STONE, content='Entry', size=14, anchor='m', margin=(0.15, 0.02, 0.08, 0.02))

    s = cs(D, C, 'Process step profile', 'processstep', N('Standard profile for one process step.', 'Duplicate per step; mark position in the chevron row.',
                                                          'The same profile for all steps makes process handbooks comparable.'))
    chevrons(s, ['Step 1', 'Step 2', 'Step 3', 'Step 4', 'Step 5'], L, 1.45, W, 0.45, active=3, size=12)
    bw = (W - 2 * 0.55) / 3
    for i, h in enumerate(['Input', 'Activities', 'Output']):
        x = L + i * (bw + 0.55)
        rect(s, x, 2.15, bw, 1.9, fill=BLUE if i == 1 else LSAND)
        label(s, x + 0.2, 2.3, bw - 0.4, h, color=WHITE if i == 1 else BLUE)
        text(s, x + 0.2, 2.6, bw - 0.4, 1.4, bl(B3, sa=4), size=14, color=WHITE if i == 1 else BLACK)
        if i < 2:
            icon(s, 'arrow_forward', x + bw + 0.14, 2.95, 0.28, TAUPE)
    mw = (W - 3 * 0.3) / 4
    for i, a in enumerate(['Owner', 'Systems', 'KPI', 'Risks']):
        x = L + i * (mw + 0.3)
        hline(s, x, 4.4, mw, BLACK); label(s, x, 4.52, mw, a, color=DGRAY); text(s, x, 4.82, mw, 0.8, '[Entry]', size=14)

    s = cs(D, C, 'Value stream', 'vsm', N('Value stream map – shows waste in lead time.', 'Enter process times (PT) and waiting times; update totals.'))
    n_ = 5; bw = 1.75; gap = (W - n_ * bw) / (n_ - 1)
    for i in range(n_):
        x = L + i * (bw + gap)
        rect(s, x, 1.7, bw, 1.15, fill=SAND, content=[{'t': 'Step %d' % (i + 1), 'size': 14}, {'t': 'PT 0.0 d', 'size': 12, 'color': BLUE}], align='c', anchor='m')
        if i < n_ - 1:
            tx = x + bw + gap / 2 - 0.3
            shape(s, SH.ISOSCELES_TRIANGLE, tx, 1.95, 0.6, 0.5, fill=STONE)
            text(s, tx - 0.3, 2.5, 1.2, 0.3, '0 d', size=13, align='c')
    yb = 3.45
    for i in range(n_):
        x = L + i * (bw + gap)
        hline(s, x, yb + 0.45, bw, BLUE, 2.5)
        if i < n_ - 1:
            hline(s, x + bw, yb, gap, BLACK, 1.5)
            vline(s, x + bw, yb, 0.45, BLACK, 1.0); vline(s, x + bw + gap, yb, 0.45, BLACK, 1.0)
    for i, (a, b) in enumerate([('0.0 days', 'Lead time'), ('0.0 days', 'Process time'), ('0%', 'Value-added share')]):
        kpi(s, L + i * 4.3, 4.45, 3.9, a, b, bigsize=28)

    s = cs(D, C, 'As-is and to-be process', 'tobeprocess', N('Comparison of current and future process.', 'Dropped steps with a cross, new steps in blue.'))
    label(s, L, 1.45, 6, 'As-is process', color=DGRAY)
    keep = {0, 2, 5, 6}
    bw = (W - 6 * 0.1) / 7
    for i in range(7):
        x = L + i * (bw + 0.1)
        rect(s, x, 1.8, bw, 0.85, fill=SAND if i in keep else WHITE, line=None if i in keep else STONE, content='Step', size=13, align='c', anchor='m',
             color=BLACK if i in keep else DGRAY)
        if i not in keep:
            icon(s, 'cancel', x + bw / 2 - 0.14, 2.75, 0.28, TAUPE)
    label(s, L, 3.35, 6, 'To-be process', color=BLUE)
    bw2 = (W - 3 * 0.45) / 4
    for i in range(4):
        x = L + i * (bw2 + 0.45); new = i in (1, 2)
        rect(s, x, 3.7, bw2, 0.95, fill=BLUE if new else SAND, content='Step', size=15, color=WHITE if new else BLACK, align='c', anchor='m')
        if i < 3:
            icon(s, 'arrow_forward', x + bw2 + 0.1, 4.04, 0.24, TAUPE)
    rect(s, L, 4.95, 0.22, 0.22, fill=BLUE); text(s, L + 0.32, 4.91, 3, 0.3, 'New / digitised', size=12)
    icon(s, 'cancel', L + 2.6, 4.93, 0.24, TAUPE); text(s, L + 2.95, 4.91, 3, 0.3, 'Dropped', size=12)
    sowhat(s, KT)

    s = cs(D, C, 'Decision logic', 'flowchart', N('Flow chart with decisions – for SOPs and to-be processes.', 'Start/end = rounded, step = rectangle, decision = diamond.'))
    def box(x, y, t, kind='step', w=2.3, h=0.75):
        if kind == 'term':
            return rrect(s, x, y, w, h, fill=BLACK, radius=0.37, content=t, size=13, color=WHITE, align='c', anchor='m')
        if kind == 'dec':
            return shape(s, SH.DIAMOND, x, y, w, h + 0.5, fill=BLUE, content=t, size=13, color=WHITE, align='c', anchor='m', margin=0.05)
        return rect(s, x, y, w, h, fill=SAND, content=t, size=13, align='c', anchor='m')
    box(L, 1.7, 'Start', 'term', 1.8); arrow(s, L + 1.8, 2.07, L + 2.3, 2.07)
    box(L + 2.3, 1.7, 'Step'); arrow(s, L + 4.6, 2.07, L + 5.1, 2.07)
    box(L + 5.1, 1.45, 'Question?', 'dec', 2.3, 0.75); arrow(s, L + 7.4, 2.07, L + 8.0, 2.07); text(s, L + 7.45, 1.72, 0.6, 0.3, 'No', size=12)
    box(L + 8.0, 1.7, 'Step'); arrow(s, L + 10.3, 2.07, L + 10.8, 2.07)
    box(L + 10.8, 1.7, 'End', 'term', 1.8)
    arrow(s, L + 6.25, 2.7, L + 6.25, 3.4); text(s, L + 6.35, 2.85, 0.6, 0.3, 'Yes', size=12)
    box(L + 5.1, 3.4, 'Step')
    arrow(s, L + 7.4, 3.77, L + 8.0, 3.77)
    box(L + 8.0, 3.4, 'Step')
    elbow(s, L + 10.3, 3.77, L + 11.7, 2.45)

    s = cs(D, C, 'KPI profile', 'processkpi', N('KPI definitions – basis for reporting and dashboards.', 'Agree definition, source and owner with the customer.',
                                                'Every KPI needs definition, data source, rhythm and owner – otherwise it is not measured.'))
    rows = [['KPI', 'Definition', 'Today', 'Target', 'Source / rhythm', 'Owner']] + [['KPI %d' % i, '[Definition]', '[0]', '[0]', '[System, rhythm]', '[Role]'] for i in range(1, 5)]
    tb = table(s, L, 1.5, W, rows, colw=[2.2, 3.6, 1.1, 1.1, 2.6, 2.0], rowh=0.7, size=14, aligns=['l', 'l', 'c', 'c', 'l', 'l'])
    for r in range(1, 5):
        fill_tf(tb.cell(r, 3).text_frame, '[0]', size=14, color=BLUE, align='c')


# =====================================================================================
def s09(D):
    C = 'References'
    divider(D, 9, 'References', N('References, case studies and testimonials.', 'Only approved references; otherwise anonymise (“Top-20 pharma company”).',
                                   'The best reference has the most similar problem – not the biggest logo.'))
    s = cs(D, C, 'Overview', 'refs', N('Overview of relevant references.', 'Pick six that match the customer; quantify results.'))
    cw, chh = (W - 2 * 0.3) / 3, 2.0
    for i in range(6):
        x = L + (i % 3) * (cw + 0.3); y = 1.5 + (i // 3) * (chh + 0.3)
        rect(s, x, y, cw, chh, fill=LSAND)
        rrect(s, x + 0.25, y + 0.25, 1.3, 0.3, fill=WHITE, line=BLUE, radius=0.15, content='Segment', size=11, color=BLUE, align='c', anchor='m', margin=0)
        text(s, x + 0.25, y + 0.75, cw - 0.5, 0.4, 'Project title', size=18)
        text(s, x + 0.25, y + 1.25, cw - 0.5, 0.5, 'Result in one line', size=15, color=BLUE)

    s = cs(D, C, 'Case study', 'case1', N('Standard case study one-pager.', 'Max. three points per block; picture via Change Picture.',
                                         'Put the result into the title.'), source='Customer: [anonymised description]')
    picture(s, IMG + 'image21.jpeg', 9.0, 1.45, R - 9.0, 4.35)
    bw = (9.0 - L - 0.3 - 2 * 0.25) / 3
    for i, h in enumerate(['Challenge', 'Approach', 'Result']):
        x = L + i * (bw + 0.25); f = BLUE if i == 2 else LSAND
        rect(s, x, 1.45, bw, 2.6, fill=f)
        label(s, x + 0.2, 1.65, bw - 0.4, h, color=WHITE if i == 2 else BLUE)
        text(s, x + 0.2, 1.98, bw - 0.4, 1.9, bl(['Point', 'Point'], sa=6), size=14, color=WHITE if i == 2 else BLACK)
    for i in range(3):
        kpi(s, L + i * (bw + 0.25), 4.35, bw, ['00%', '0 weeks', '€0.0m'][i], 'Result KPI', bigsize=28)

    s = cs(D, C, 'Case study with results', 'case2', N('Case study with before/after figures.', 'Edit chart data.'), source='Customer: [anonymised description]')
    bar(s, L, 1.4, 7.4, 4.4, ['KPI 1', 'KPI 2', 'KPI 3'], [('Before', [0.5, 0.7, 0.8]), ('After', [0.65, 0.8, 0.9])],
        fmt='0%', colors=[STONE, BLUE], gap=80, overlap=-10, legend='b', size=13)
    xr = 8.2; wr = R - xr
    label(s, xr, 1.5, wr, 'What made the difference')
    for i in range(3):
        y = 1.9 + i * 1.25
        num(s, xr, y, i + 1, 0.42)
        text(s, xr + 0.6, y - 0.02, wr - 0.6, 0.35, 'Success factor', size=17)
        text(s, xr + 0.6, y + 0.4, wr - 0.6, 0.6, D1, size=14, color=DGRAY)

    s = cs(D, C, 'Customer logos', 'logos', N('Logo wall for first meetings.', 'Insert logos in greyscale; only with written approval.'), sticker_='Approval needed')
    cw, chh = (W - 3 * 0.3) / 4, 1.2
    for i in range(12):
        x = L + (i % 4) * (cw + 0.3); y = 1.6 + (i // 4) * (chh + 0.3)
        rect(s, x, y, cw, chh, fill=LSAND, content='[Logo]', size=14, color=TAUPE, align='c', anchor='m')


# =====================================================================================
def s10(D):
    C = 'Proposal'
    divider(D, 10, 'Proposal', N('Proposal presentation: scope, packages, fees, payment plan, mutual action plan.', 'Align figures with the written proposal.',
                                 'Never show the price without the value – present the business case first.'))
    s = cs(D, C, 'Proposal at a glance', 'offer', N('Executive summary of the proposal.', 'Take all four fields from the written proposal.'),
           source='Details: proposal no. [xxx], [date]')
    cw = (W - 0.3) / 2; chh = 1.95
    for i, (h, ic) in enumerate([('Your situation', 'search'), ('Our scope', 'checklist'), ('Your result', 'flag'), ('Your investment', 'euro')]):
        x = L + (i % 2) * (cw + 0.3); y = 1.5 + (i // 2) * (chh + 0.3)
        f = BLUE if i == 2 else LSAND
        rect(s, x, y, cw, chh, fill=f)
        icon(s, ic, x + 0.3, y + 0.3, 0.45, WHITE if f == BLUE else BLUE)
        text(s, x + 1.0, y + 0.3, cw - 1.3, 0.4, h, size=19, color=WHITE if f == BLUE else BLACK)
        text(s, x + 1.0, y + 0.85, cw - 1.3, 0.9, D1, size=15, color=WHITE if f == BLUE else BLACK)

    s = cs(D, C, 'Scope', 'scope', N('In/out of scope – prevents misunderstandings.', 'Align with the service description.', 'Address out-of-scope actively – it builds trust.'))
    cw = (W - 0.5) / 2
    for i, (h, ic, col) in enumerate([('In scope', 'check_circle', BLUE), ('Out of scope', 'cancel', TAUPE)]):
        x = L + i * (cw + 0.5)
        rect(s, x, 1.5, cw, 0.7, fill=BLUE if i == 0 else SAND, content=h, size=17, color=WHITE if i == 0 else BLACK, anchor='m', margin=(0.25, 0, 0.1, 0))
        for j in range(4):
            y = 2.4 + j * 0.7
            icon(s, ic, x + 0.05, y + 0.08, 0.36, col)
            text(s, x + 0.6, y, cw - 0.6, 0.52, 'Service', size=15, anchor='m')
            hline(s, x, y + 0.62, cw, STONE, 0.5)

    s = cs(D, C, 'Packages', 'packages', N('Good-better-best packages – moves the discussion from “if” to “which”.', 'Recommended package in the middle.'))
    x0 = L + 4.0; cw = (R - x0 - 2 * 0.2) / 3
    inc = [[1, 1, 0, 0, 0], [1, 1, 1, 1, 0], [1, 1, 1, 1, 1]]
    for j in range(3):
        x = x0 + j * (cw + 0.2); hi = j == 1
        rect(s, x, 1.45, cw, 4.35, fill=LBLUE if hi else LSAND)
        rect(s, x, 1.45, cw, 1.2, fill=BLUE if hi else SAND)
        if hi:
            label(s, x + 0.2, 1.55, cw - 0.4, 'Recommended', color=WHITE)
        text(s, x + 0.2, 1.8, cw - 0.4, 0.4, 'Package %d' % (j + 1), size=19, color=WHITE if hi else BLACK)
        text(s, x + 0.2, 2.2, cw - 0.4, 0.4, '€[price]', size=16, color=WHITE if hi else BLUE)
        for i, v in enumerate(inc[j]):
            y = 2.9 + i * 0.55
            if v:
                icon(s, 'check_circle', x + cw / 2 - 0.16, y + 0.08, 0.32, BLUE)
            else:
                hline(s, x + cw / 2 - 0.1, y + 0.24, 0.2, STONE, 1)
    for i in range(5):
        y = 2.9 + i * 0.55
        text(s, L, y, 3.9, 0.48, 'Service %d' % (i + 1), size=15, anchor='m')
        hline(s, L, y + 0.52, R - L, STONE, 0.5)

    s = cs(D, C, 'Fees', 'fees', N('Fee overview per phase.', 'Copy values from the calculation; check totals.'))
    rows = [['Phase', 'Services', 'Days', 'Fee', 'Billing']] + [['Phase %d' % i, '[Services]', '[0]', '€[0]k', '[Fixed / T&M]'] for i in range(1, 5)] + \
           [[('Total', {'color': BLUE}), '', ('[0]', {'color': BLUE}), ('€[0]k', {'color': BLUE}), '']]
    table(s, L, 1.5, W, rows, colw=[2.4, 4.6, 1.6, 2.0, 2.0], rowh=0.62, size=15, aligns=['l', 'l', 'r', 'r', 'l'])
    sowhat(s, KT)

    s = cs(D, C, 'Payment plan', 'payment', N('Milestone-based payment plan – reduces perceived risk.', 'Align milestones and shares with the contract.'))
    cw = W / 5; y0 = 3.1
    hline(s, L, y0, W, BLACK, 1.25)
    for i in range(5):
        x = L + i * cw
        oval(s, x, y0 - 0.2, 0.4, fill=BLUE if i else BLACK)
        text(s, x, 1.6, cw - 0.2, 0.7, '00%', size=36, color=BLUE)
        text(s, x, 2.4, cw - 0.2, 0.3, 'Milestone %d' % i if i else 'Order', size=16)
        text(s, x, y0 + 0.4, cw - 0.25, 0.6, 'Trigger for payment', size=14)
    rect(s, L, 4.5, W, 1.3, fill=LSAND)
    text(s, L + 0.3, 4.6, W - 0.6, 1.1, bl(['Payment terms', 'Billing of time & material', 'Travel expenses'], sa=4), size=15)

    s = cs(D, C, 'Mutual action plan', 'map_plan', N('Joint plan until signature – makes the decision process transparent.', 'Agree steps with the customer; mirror them in the CRM.',
                                                     'A mutual action plan shortens sales cycles because delays become visible immediately.'), source='As of: [date]')
    rows = [['Step', 'Owner', 'Date', 'Status']] + [['Step %d' % i, '[Customer / Körber]', '[DD.MM.]', ''] for i in range(1, 7)]
    table(s, L, 1.5, W, rows, colw=[5.2, 3.8, 1.6, 2.0], rowh=0.58, size=15)
    st = ['g', 'g', 'y', 'n', 'n', 'n']
    for r in range(1, 7):
        rag(s, L + W * (10.6 / 12.6) + 0.15, 1.5 + r * 0.58 + 0.19, st[r - 1])
        text(s, L + W * (10.6 / 12.6) + 0.45, 1.5 + r * 0.58, 1.5, 0.58, {'g': 'Done', 'y': 'In progress', 'n': 'Open'}[st[r - 1]], size=13, anchor='m')
